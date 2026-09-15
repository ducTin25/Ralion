from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "PM" / "Ralion_Project_Report.docx"


def set_cell_border(cell, color="DADCE0", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def configure_table(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            margins = tc_pr.find(qn("w:tcMar"))
            if margins is None:
                margins = OxmlElement("w:tcMar")
                tc_pr.append(margins)
            for side, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                node = margins.find(qn("w:" + side))
                if node is None:
                    node = OxmlElement("w:" + side)
                    margins.append(node)
                node.set(qn("w:w"), str(value))
                node.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, value in enumerate(headers):
        p = table.rows[0].cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(value)
        run.bold = True
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.add_run(str(value))
    configure_table(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def setup_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    normal_ppr = normal.element.get_or_add_pPr()
    normal_border = normal_ppr.find(qn("w:pBdr"))
    if normal_border is not None:
        normal_ppr.remove(normal_border)

    specs = {
        "Heading 1": (20, "000000", 20, 6),
        "Heading 2": (16, "000000", 18, 6),
        "Heading 3": (14, "434343", 16, 4),
    }
    for name, (size, color, before, after) in specs.items():
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15


def p(doc, text, bold_prefix=None):
    para = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        para.add_run(bold_prefix).bold = True
        para.add_run(text[len(bold_prefix):])
    else:
        para.add_run(text)
    return para


def build():
    doc = Document()
    setup_styles(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(3)
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    run = title.add_run("RALION")
    run.font.name = "Arial"
    run.font.size = Pt(26)
    run.font.bold = False
    run.font.color.rgb = RGBColor(0, 0, 0)
    subtitle = doc.add_paragraph("AI Technical Onboarding Buddy - Project Report")
    subtitle.runs[0].font.size = Pt(16)
    meta = doc.add_paragraph("Product, business, AI evaluation and technical report | Version 1.0 | 24 August 2026")
    meta.runs[0].font.color.rgb = RGBColor.from_string("555555")

    doc.add_heading("Executive summary", level=1)
    p(doc, "Ralion supports software engineers when they join an active software project for the first time. It brings fragmented project and company knowledge into one governed onboarding journey: approved documents, a reusable master template, an AI-assisted candidate plan, task-level guidance, grounded chat with citations, progress tracking, blockers and the path to a first contribution.")
    p(doc, "The product uses AI where synthesis and discovery create leverage: document classification, onboarding guidance generation, retrieval-grounded answers and engineering-convention discovery from GitHub pull-request review comments. Human approval remains the control point. Project Managers approve project knowledge, templates, plans and proposed conventions before they affect engineers.")
    p(doc, "The current evaluation compares AI-generated task guidance with deterministic Baseline B0 across 30 golden cases. AI scores higher than B0 in 70% of cases and achieves an overall judge score of approximately 4.5/5 versus approximately 4.0/5. The strongest results are in Architecture, Codebase and Setup. Current observed plan-generation economics are approximately USD 0.04 per plan and approximately 45 seconds of latency; these are operational observations, not an SLA.")

    doc.add_heading("1. Problem and opportunity", level=1)
    p(doc, "Joining an established project is not a documentation problem alone. Knowledge is distributed across repositories, README files, architecture decisions, policies, pull requests, review comments and people. New engineers must determine what is current, what is relevant to their role, what they are allowed to access and what to do next.")
    add_bullets(doc, [
        "Engineers spend time searching, waiting for access and repeating questions that have already been answered.",
        "Project Managers and senior engineers repeatedly reconstruct the same onboarding path for each new member.",
        "Static checklists become stale and rarely explain why a task matters or where its evidence comes from.",
        "Progress and blockers are difficult to observe until they already affect delivery.",
        "Important engineering conventions often remain implicit inside code-review history.",
    ])
    p(doc, "Ralion addresses this gap without replacing Jira, a wiki, GitHub or a mentor. It connects those sources into an ordered, permission-aware and reviewable journey.")

    doc.add_heading("2. Product scope and value proposition", level=1)
    add_table(doc, ["Stakeholder", "Value delivered"], [
        ["Engineer", "A clear checklist, task guidance, cited answers, visible dependencies and a direct blocker workflow."],
        ["Project Manager", "Reusable templates, AI-assisted plan creation, controlled document ingestion, approvals and progress visibility."],
        ["Senior/Tech Lead", "Fewer repeated explanations and more consistent guidance before the first pull request."],
        ["Organization", "A governed onboarding process with measurable time, quality and support-cost outcomes."],
    ], [2100, 7260])
    p(doc, "In scope for the MVP: authentication and role-based access, project membership, policy and project documents, document versioning and chunking, master templates, AI-assisted plans, engineer checklists, RAG chat, citations, blockers, First Task/First PR tracking, GitHub document synchronization and convention review.")
    p(doc, "Out of scope: automatic permission granting, autonomous code changes or merges, employee performance scoring, replacement of authoritative source systems and unsupervised publication of AI-generated rules.")

    doc.add_heading("3. Actors and responsibilities", level=1)
    add_table(doc, ["Role", "Primary responsibility", "AI interaction"], [
        ["Admin", "Manage users, projects, membership and role boundaries.", "Enables governed access; does not approve AI content."],
        ["HR", "Manage company policies and their publication lifecycle.", "Provides approved policy sources used by plans and chat."],
        ["Project Manager", "Prepare project knowledge, manage master templates, create/review/release onboarding plans and review conventions.", "Controls AI inputs and approves AI outputs before release."],
        ["Engineer", "Complete the checklist, use cited chat, acknowledge policies, report blockers and reach the first contribution.", "Consumes approved AI-assisted guidance within access scope."],
    ], [1500, 4560, 3300])

    doc.add_heading("4. End-to-end onboarding journey", level=1)
    add_numbered(doc, [
        "Admin creates the project, users and memberships with explicit roles.",
        "HR publishes company policies; the PM uploads project documents or connects a GitHub repository.",
        "The knowledge pipeline versions, parses, chunks and embeds active documents for retrieval.",
        "The PM maintains an approved Master Template containing ordered task definitions and dependencies.",
        "AI combines the approved template with eligible project and policy sources to draft task guidance and citations.",
        "The PM reviews, edits, regenerates where necessary, then approves and releases the plan to an engineer.",
        "The engineer completes tasks, asks grounded questions, acknowledges policies and raises blockers.",
        "The PM and supporting owners track progress, resolve blockers and close onboarding after the First PR criteria are met.",
    ])

    doc.add_heading("5. AI capabilities and governance", level=1)
    add_table(doc, ["Capability", "Mechanism", "Control"], [
        ["Document classification", "AI proposes the most suitable knowledge category when metadata is uncertain.", "PM confirms or corrects classification before use."],
        ["Plan guidance generation", "A three-phase pipeline retrieves eligible sources, generates structured task guidance and verifies quoted evidence.", "Candidate content is editable and must be approved before release."],
        ["Grounded RAG Chat", "Hybrid retrieval combines lexical and vector evidence, then the LLM answers within project and policy scope.", "Answers carry citations; retrieval and visibility obey active-version and access rules."],
        ["Convention discovery", "Merged PR review comments are collected, normalized, deduplicated and clustered into proposed engineering conventions.", "PM reviews proposals; only approved conventions are shown to engineers and indexed for retrieval."],
    ], [2100, 4380, 2880])
    p(doc, "Deterministic controls remain outside the LLM: role checks, document status, version selection, template structure, plan state transitions, quote verification, blocker ownership and approval gates. This separation reduces the blast radius of model errors.")

    doc.add_heading("6. GitHub integration", level=1)
    doc.add_heading("6.1 Repository document synchronization", level=2)
    p(doc, "The PM enters an owner/repository identifier and branch. The system reads supported documentation files from GitHub, creates project Knowledge Documents, versions the content, generates chunks and embeddings, and makes approved active content available to plan generation and RAG Chat. Repository access is limited by the configured GitHub token and repository permissions; a token does not grant universal access to every repository.")
    doc.add_heading("6.2 Pull-request review convention discovery", level=2)
    p(doc, "A separate pipeline scans merged pull requests and their review comments. It converts recurring review feedback into evidence-backed convention candidates, groups duplicates and sends them to PM Rule Review. Approval turns a proposal into a project convention visible to engineers and available to retrieval. Rejected or pending proposals do not become project guidance.")
    p(doc, "These two GitHub flows solve different problems: repository sync supplies formal project knowledge; review mining captures working practices that may never have been documented.")

    doc.add_heading("7. Technical architecture", level=1)
    add_table(doc, ["Layer", "Current implementation"], [
        ["Frontend", "Next.js 16, React 19, TypeScript, TanStack Query, next-intl; role-specific PM and Engineer workspaces."],
        ["API", "FastAPI application with versioned routes, request trace IDs, health/readiness endpoints and role-aware use cases."],
        ["Application/domain", "Services and use cases for knowledge, templates, plans, chat, memberships, blockers and conventions."],
        ["Persistence", "PostgreSQL 16 with SQLAlchemy/Alembic; pgvector for embeddings and pg_search/ParadeDB for lexical retrieval."],
        ["AI providers", "Configurable LLM and embedding providers; DeepSeek used for evaluated content generation and a separate Groq-hosted judge for evaluation."],
        ["Operations", "Docker Compose for backend and database, health/readiness probes, structured logging and optional Langfuse telemetry."],
    ], [1900, 7460])
    p(doc, "Core knowledge flow: source file -> KnowledgeDocument -> immutable DocumentVersion -> parsed chunks -> embeddings/indexes -> filtered retrieval -> cited AI output. Only ACTIVE documents and ACTIVE versions are eligible for production retrieval.")

    doc.add_heading("8. Security, privacy and guardrails", level=1)
    add_bullets(doc, [
        "Least-privilege, project-scoped access with explicit Admin, HR, PM and Engineer role checks.",
        "No automatic access grants, code changes, merges or employee-performance decisions.",
        "Active-version filtering prevents obsolete document versions from entering retrieval.",
        "Quote verification rejects unsupported verbatim evidence while preserving valid citations.",
        "Trace IDs and structured error responses support diagnosis without exposing provider secrets.",
        "API keys and GitHub tokens belong in deployment secrets, never in source control or reports.",
        "Human approval is required before plans and discovered conventions become authoritative guidance.",
    ])

    doc.add_heading("9. Evaluation design", level=1)
    p(doc, "The evaluation tests whether AI-generated onboarding task guidance improves on the existing deterministic behavior. It uses 30 golden cases across four projects and six task categories. All 30 cases completed successfully in the recorded run.")
    add_table(doc, ["Component", "Definition"], [
        ["Baseline B0", "Deterministic copy of the PM-authored instruction_template. It does not retrieve documents and does not call an LLM."],
        ["AI branch", "The production content-generation path using retrieved project/policy context and DeepSeek."],
        ["Independent judge", "Groq openai/gpt-oss-120b, separate from the generator to reduce self-preference bias."],
        ["Run protocol", "One AI generation per case (runs=1); total recorded evaluation time 2,198.5 seconds."],
        ["Golden set", "30 cases: Access 8, Architecture 4, Codebase 4, Company 2, Orientation 4 and Setup 8."],
    ], [2100, 7260])

    doc.add_heading("10. Metrics", level=1)
    add_table(doc, ["Metric", "What it measures", "Interpretation"], [
        ["Faithfulness", "Share of generated claims supported by retrieved source text.", "Higher is better; low scores may signal hallucination or an evaluation-context artifact."],
        ["Context Precision", "Share of retrieved chunks relevant to the task objective.", "Higher means less retrieval noise; safety-oriented broad retrieval can lower it intentionally."],
        ["Context Recall", "Share of reference-answer ideas present in retrieved context.", "Computed only for cases with reference content."],
        ["LLM Judge", "Correctness, relevance, completeness and coherence, each scored 1-5 and averaged.", "Measures end-to-end usefulness of the generated guidance."],
        ["AI win rate", "Share of cases where Judge AI is strictly greater than Judge B0.", "Direct per-case comparison; ties do not count as wins."],
    ], [1800, 4560, 3000])

    doc.add_heading("11. Evaluation results", level=1)
    add_table(doc, ["Category", "N", "Faith AI", "Faith B0", "Ctx precision", "Ctx recall", "Judge AI", "Judge B0", "AI wins"], [
        ["Access", 8, "0.97 +/- 0.04", "0.62", "0.29", "1.00 (N=2)", "4.38 +/- 0.40", "4.28", "25%"],
        ["Architecture", 4, "1.00 +/- 0.00", "0.75", "0.87", "1.00 (N=2)", "4.88 +/- 0.12", "3.94", "100%"],
        ["Codebase", 4, "0.97 +/- 0.03", "0.75", "0.79", "1.00 (N=2)", "4.69 +/- 0.32", "3.56", "75%"],
        ["Company", 2, "0.74 +/- 0.26", "0.17", "0.89", "1.00 (N=2)", "3.75 +/- 0.25", "3.88", "50%"],
        ["Orientation", 4, "1.00 +/- 0.00", "0.71", "0.68", "1.00 (N=2)", "4.81 +/- 0.11", "4.44", "75%"],
        ["Setup", 8, "1.00 +/- 0.01", "0.88", "0.85", "1.00 (N=2)", "4.59 +/- 0.12", "4.06", "100%"],
    ], [1250, 450, 1200, 900, 1050, 1050, 1250, 900, 1310])
    p(doc, "Source: docs/PM/evaluation/eval-report.md, generated by scripts/eval_plan_content.py.")
    add_bullets(doc, [
        "Overall AI win rate: 70% across 30 cases.",
        "Overall judge score: approximately 4.5/5 for AI versus approximately 4.0/5 for B0.",
        "Strongest evidence of value: Architecture and Setup (100% win rate), followed by Codebase and Orientation (75%).",
        "Access has high faithfulness but low context precision because the retrieval design intentionally includes broad security policy context.",
        "Company faithfulness is affected by judge-context truncation: generation saw up to 70 chunks while the judge saw a capped subset.",
    ])

    doc.add_heading("12. Cost and latency", level=1)
    add_table(doc, ["Operational measure", "Observed value", "How to report it"], [
        ["Plan-generation cost", "Approximately USD 0.04 per plan", "Current observation; varies with model pricing, prompt size, retrieved context and number of tasks."],
        ["Plan-generation latency", "Approximately 45 seconds per plan", "Current observation; varies with provider load, network, document volume and retry behavior."],
        ["Evaluation runtime", "2,198.5 seconds for 30 cases", "Recorded evaluation run with one generation per case and additional judging calls."],
    ], [2500, 2100, 4760])
    p(doc, "For 100 generated plans at the same observed unit cost, direct model cost would be approximately USD 4.00. This is a linear planning estimate, not a production invoice forecast. Production monitoring should track p50/p95 latency, input/output tokens, retries, failure rate and cost by project and plan version.")

    doc.add_heading("13. Limitations and risks", level=1)
    add_bullets(doc, [
        "The evaluation ran once per case, so it does not measure generation stability across repeated runs.",
        "Only 12 cases contain reference_content for Context Recall; Company has only two cases.",
        "The judge context is capped, creating a known measurement artifact for large policy corpora.",
        "LLM-as-Judge results have not yet been calibrated against at least ten PM-scored cases.",
        "The 70% AI win rate is below an 80% aspirational straight-through approval target; PM review remains necessary.",
        "GitHub synchronization quality depends on token permissions, repository structure, file support and API rate limits.",
        "Operational cost and latency are preliminary observations and require production-grade instrumentation before SLA or ROI commitments.",
    ])

    doc.add_heading("14. Recommended improvement plan", level=1)
    add_numbered(doc, [
        "Fix judge-context sampling so large categories use representative chunks instead of the first fixed subset.",
        "Separate Context Precision for project documents and company policies to make safety-related retrieval trade-offs visible.",
        "Repeat each golden case at least three times and report within-case variance.",
        "Have PMs manually score at least ten cases and measure agreement with the automated judge.",
        "Expand Company and Access cases, including negative/no-evidence and permission-boundary scenarios.",
        "Instrument plan generation with p50/p95 latency, token usage, cost, retry count, provider error rate and approval/edit rate.",
        "Measure real pilot outcomes: time to First PR, senior support hours, blocker resolution time and first-PR rework.",
    ])

    doc.add_heading("15. Pilot measurement framework", level=1)
    add_table(doc, ["Outcome", "Baseline", "Pilot measurement", "Decision use"], [
        ["Time to First PR", "Median days for comparable recent hires/transfers.", "Median and distribution for Ralion users.", "Primary outcome and North Star proxy."],
        ["Senior support effort", "Hours/week spent on onboarding questions and setup.", "Self-reported plus tagged support/blocker activity.", "Tests operational savings."],
        ["Plan acceptance", "Not applicable for B0.", "% AI tasks approved unchanged; edit distance and regeneration rate.", "Tests content usefulness."],
        ["Answer quality", "Current manual search/mentor response.", "Grounded-answer success, citation validity and user feedback.", "Tests trust and knowledge access."],
        ["First PR quality", "Review cycles and rework for comparable contributors.", "Review cycles, convention violations and time to merge.", "Tests downstream contribution quality."],
    ], [1600, 2500, 2860, 2400])

    doc.add_heading("16. Go / pivot / stop criteria", level=1)
    add_bullets(doc, [
        "GO: reliable access controls and citations, measurable reduction in onboarding friction, acceptable PM approval effort and no critical safety failures.",
        "PIVOT: AI guidance adds limited value over B0 in specific categories; keep deterministic templates and redesign retrieval/prompts for those categories.",
        "STOP: permission leakage, unsupported authoritative claims, uncontrolled publication, or no measurable improvement in engineer and PM outcomes after a representative pilot.",
    ])

    doc.add_heading("17. Conclusion", level=1)
    p(doc, "Ralion is best understood as a governed onboarding operating layer, not a general chatbot. It converts approved company and project knowledge into a reusable plan, supports execution with evidence and exposes progress and blockers to the people responsible for resolution. The evaluation provides credible early evidence that AI improves task guidance over deterministic template copy, especially for Architecture, Codebase and Setup. The next stage should focus on measurement quality, human calibration and real pilot outcomes before making enterprise-scale ROI or SLA claims.")

    doc.add_heading("Appendix A. Evidence base", level=1)
    add_bullets(doc, [
        "docs/PM/BO_06_CLAUDE_PROJECT_SOURCE_OF_TRUTH.md - approved product, business, UX and technical scope.",
        "docs/PM/evaluation/eval-report.md - generated 30-case evaluation results.",
        "docs/PM/evaluation/plan-evaluation.md - evaluation methodology and baseline design.",
        "docs/PM/evaluation/report-evaluation.md - implemented evaluation pipeline and judge-provider decisions.",
        "docs/PM/Phase-4/eval-golden-dataset.jsonl - 30-case golden dataset.",
        "src/main.py, src/modules, src/api/routers, docker-compose.yml and frontend/package.json - current implementation evidence.",
    ])

    doc.core_properties.title = "Ralion - AI Technical Onboarding Buddy Project Report"
    doc.core_properties.subject = "Product, business, AI evaluation and technical report"
    doc.core_properties.author = "Ralion Project Team"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
